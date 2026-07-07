import json
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models import ExpectedItem, ImportBatch, ImportRow, ImportStatus, Project, User, WarehouseCode, utc_now
from app.services.audit_service import audit

HEADER_MAP = {
    "код проекта": "project_code",
    "проект": "project_code",
    "project": "project_code",
    "project_code": "project_code",
    "номенклатура": "item_name",
    "товар": "item_name",
    "наименование": "item_name",
    "item": "item_name",
    "name": "item_name",
    "количество": "quantity",
    "кол-во": "quantity",
    "qty": "quantity",
    "quantity": "quantity",
    "ед. изм": "unit",
    "единица": "unit",
    "unit": "unit",
    "склад": "warehouse",
    "warehouse": "warehouse",
    "документ": "document_number",
    "номер документа": "document_number",
    "document": "document_number",
    "код 1с": "one_c_code",
    "код номенклатуры": "one_c_code",
    "one_c_code": "one_c_code",
}


def normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace("\n", " ")


def parse_warehouse(value: Any) -> WarehouseCode | None:
    text = str(value or "").strip().lower()
    if not text:
        return None
    if "аст" in text or text in {"astana", "астана"}:
        return WarehouseCode.ASTANA
    if "алм" in text or text in {"almaty", "алматы"}:
        return WarehouseCode.ALMATY
    return None


def parse_quantity(value: Any) -> int:
    if value is None or value == "":
        return 1
    try:
        return max(int(float(str(value).replace(",", "."))), 1)
    except Exception:
        return 1


def create_import_batch(db: Session, user: User, filename: str, file_bytes: bytes) -> ImportBatch:
    try:
        wb = load_workbook(filename=BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Cannot read Excel file: {exc}")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Excel is empty")

    headers = [HEADER_MAP.get(normalize_header(v), normalize_header(v)) for v in rows[0]]
    required = {"project_code", "item_name"}
    if not required.issubset(set(headers)):
        raise HTTPException(
            status_code=400,
            detail="Excel must contain columns: Код проекта and Номенклатура/Наименование",
        )

    batch = ImportBatch(filename=filename, uploaded_by_id=user.id)
    db.add(batch)
    db.flush()

    for idx, row in enumerate(rows[1:], start=2):
        data = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        project_code = str(data.get("project_code") or "").strip()
        item_name = str(data.get("item_name") or "").strip()
        if not project_code or not item_name:
            continue
        import_row = ImportRow(
            batch_id=batch.id,
            row_number=idx,
            project_code=project_code,
            item_name=item_name,
            quantity=parse_quantity(data.get("quantity")),
            unit=str(data.get("unit") or "шт").strip() or "шт",
            warehouse=parse_warehouse(data.get("warehouse")),
            document_number=str(data.get("document_number") or "").strip() or None,
            one_c_code=str(data.get("one_c_code") or "").strip() or None,
            raw_json=json.dumps({str(k): str(v) for k, v in data.items()}, ensure_ascii=False),
        )
        db.add(import_row)

    audit(db, user, "IMPORT_UPLOADED", "ImportBatch", batch.id, f"Uploaded {filename}")
    db.commit()
    db.refresh(batch)
    return batch


def confirm_import_batch(db: Session, user: User, batch_id: str):
    batch = db.get(ImportBatch, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="Import batch not found")
    if batch.status == ImportStatus.CONFIRMED:
        raise HTTPException(status_code=400, detail="Import already confirmed")

    projects_created = 0
    items_created = 0
    items_skipped = 0

    rows = db.execute(select(ImportRow).where(ImportRow.batch_id == batch.id).order_by(ImportRow.row_number)).scalars().all()
    for row in rows:
        project = db.execute(select(Project).where(Project.project_code == row.project_code)).scalar_one_or_none()
        if not project:
            project = Project(project_code=row.project_code, name=row.project_code, warehouse=row.warehouse)
            db.add(project)
            db.flush()
            projects_created += 1

        existing = db.execute(
            select(ExpectedItem).where(
                ExpectedItem.project_id == project.id,
                ExpectedItem.one_c_code == row.one_c_code,
                ExpectedItem.name == row.item_name,
                ExpectedItem.document_number == row.document_number,
            )
        ).scalar_one_or_none()
        if existing:
            items_skipped += 1
            continue

        db.add(ExpectedItem(
            project_id=project.id,
            one_c_code=row.one_c_code,
            name=row.item_name,
            quantity=row.quantity,
            unit=row.unit,
            warehouse=row.warehouse,
            document_number=row.document_number,
            source_import_id=batch.id,
        ))
        items_created += 1

    batch.status = ImportStatus.CONFIRMED
    batch.confirmed_at = utc_now()
    audit(db, user, "IMPORT_CONFIRMED", "ImportBatch", batch.id, f"created={items_created}, skipped={items_skipped}")
    db.commit()
    db.refresh(batch)
    return batch, projects_created, items_created, items_skipped
